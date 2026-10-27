# Gera public/downloads/planilha-cmv-ficha-tecnica-tanocaixa.xlsx
# Uso: python blog/recursos-src/planilha-cmv.py   (precisa de openpyxl)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

MAR, SOL, CREME, TRILHO = '2E5F73', 'EFAB5C', 'F7F5EA', 'E1DFCE'
cab = Font(bold=True, color='FFFFFF', name='Arial')
f_cab = PatternFill('solid', fgColor=MAR)
f_in = PatternFill('solid', fgColor='FFF4DC')   # célula para preencher
f_res = PatternFill('solid', fgColor=SOL)
borda = Border(*(Side(style='thin', color=TRILHO),) * 4)
BRL = 'R$ #,##0.00'
PCT = '0.0%'

wb = Workbook()

# ---------- Ficha técnica ----------
ws = wb.active
ws.title = 'Ficha técnica'
ws['A1'] = 'Ficha técnica do prato'
ws['A1'].font = Font(bold=True, size=16, color=MAR, name='Arial')
ws['A2'] = 'Preencha só as células amarelas. Exemplo: pizza grande de calabresa (valores ilustrativos: troque pelos seus).'
ws['A2'].font = Font(italic=True, color='6A7A7E', name='Arial')

campos = [('Prato', 'Pizza grande (8 fatias) de calabresa', None),
          ('Preço de venda (R$)', 59.90, BRL),
          ('CMV que você quer atingir (%)', 0.30, PCT),
          ('Embalagem e descartáveis (R$)', 2.50, BRL)]
for i, (r, v, fmt) in enumerate(campos, start=4):
    ws.cell(i, 1, r).font = Font(bold=True, name='Arial')
    c = ws.cell(i, 2, v); c.fill = f_in; c.border = borda
    if fmt: c.number_format = fmt

heads = ['Ingrediente', 'Quantidade usada no prato', 'Unidade', 'Preço de compra (R$)', 'Quantidade comprada', 'Perda / aparas (%)', 'Custo no prato (R$)']
for j, h in enumerate(heads, start=1):
    c = ws.cell(9, j, h); c.font = cab; c.fill = f_cab; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = borda
ws.row_dimensions[9].height = 34

ex = [('Massa (farinha, fermento, óleo)', 1, 'un', 3.20, 1, 0.00),
      ('Molho de tomate', 1, 'un', 1.60, 1, 0.00),
      ('Mussarela', 300, 'g', 38.00, 1000, 0.03),
      ('Calabresa', 80, 'g', 35.00, 1000, 0.05),
      ('Cebola, azeitona e orégano', 1, 'un', 1.20, 1, 0.10)]
for i in range(10, 25):
    k = i - 10
    for j in range(1, 7):
        c = ws.cell(i, j); c.fill = f_in; c.border = borda
        if k < len(ex): c.value = ex[k][j - 1]
    ws.cell(i, 4).number_format = BRL
    ws.cell(i, 6).number_format = PCT
    g = ws.cell(i, 7, f'=IF(OR(B{i}="",D{i}="",E{i}=""),0,(B{i}*D{i}/E{i})/(1-MIN(0.95,N(F{i}))))')
    g.number_format = BRL; g.border = borda

ws['F26'] = 'Custo dos ingredientes'; ws['G26'] = '=SUM(G10:G24)'
ws['F27'] = 'Custo total do prato'; ws['G27'] = '=G26+B7'
ws['F28'] = 'CMV do prato'; ws['G28'] = '=IF(B5>0,G27/B5,0)'
ws['F29'] = 'Preço sugerido para o CMV alvo'; ws['G29'] = '=IF(B6>0,G27/B6,0)'
ws['F30'] = 'Sobra por prato (no preço atual)'; ws['G30'] = '=B5-G27'
for r in range(26, 31):
    ws.cell(r, 6).font = Font(bold=True, name='Arial'); ws.cell(r, 6).alignment = Alignment(horizontal='right')
    ws.cell(r, 7).fill = f_res if r in (28, 29) else PatternFill('solid', fgColor=CREME)
    ws.cell(r, 7).font = Font(bold=True, name='Arial'); ws.cell(r, 7).border = borda
    ws.cell(r, 7).number_format = PCT if r == 28 else BRL
ws['A32'] = 'Como ler: o custo no prato = (quantidade usada × preço ÷ quantidade comprada) ÷ (1 − perda). Perda é o que sobra no lixo (aparas, osso, pele, queima).'
ws['A33'] = 'O CMV que você quer atingir é uma meta sua. Muita gente usa de 28% a 35% como referência, mas varia por cozinha, preço e perda.'
for r in (32, 33): ws.cell(r, 1).font = Font(italic=True, color='45585D', name='Arial')
for col, w in zip('ABCDEFG', (38, 18, 10, 18, 18, 14, 20)): ws.column_dimensions[col].width = w
ws.column_dimensions['F'].width = 30
ws.freeze_panes = 'A10'

# ---------- CMV do mês ----------
w2 = wb.create_sheet('CMV do mês')
w2['A1'] = 'CMV do mês'; w2['A1'].font = Font(bold=True, size=16, color=MAR, name='Arial')
w2['A2'] = 'CMV = estoque inicial + compras − estoque final. Valores abaixo são ilustrativos.'
w2['A2'].font = Font(italic=True, color='6A7A7E', name='Arial')
linhas = [('Estoque inicial do mês (R$)', 8000), ('Compras do mês (R$)', 15500), ('Estoque final do mês (R$)', 7500), ('Faturamento do mês (R$)', 50000)]
for i, (r, v) in enumerate(linhas, start=4):
    w2.cell(i, 1, r).font = Font(bold=True, name='Arial')
    c = w2.cell(i, 2, v); c.fill = f_in; c.number_format = BRL; c.border = borda
w2['A9'] = 'CMV do mês (R$)'; w2['B9'] = '=B4+B5-B6'
w2['A10'] = 'CMV do mês (%)'; w2['B10'] = '=IF(B7>0,B9/B7,0)'
w2['A11'] = 'Sobra depois do CMV (R$)'; w2['B11'] = '=B7-B9'
for r, fmt in ((9, BRL), (10, PCT), (11, BRL)):
    w2.cell(r, 1).font = Font(bold=True, name='Arial')
    c = w2.cell(r, 2); c.fill = f_res; c.number_format = fmt; c.font = Font(bold=True, name='Arial'); c.border = borda
w2['A13'] = 'Dica: conte o estoque no último dia do mês, sempre no mesmo horário. O estoque final de um mês é o inicial do próximo.'
w2['A13'].font = Font(italic=True, color='45585D', name='Arial')
w2.column_dimensions['A'].width = 34; w2.column_dimensions['B'].width = 20

# ---------- Leia-me ----------
w3 = wb.create_sheet('Leia-me')
txt = ['Planilha de CMV e ficha técnica — Tá no Caixa (tanocaixa.com)',
       '',
       '1) Aba "Ficha técnica": uma por prato. Duplique a aba para cada prato do cardápio.',
       '2) Aba "CMV do mês": preencha no fechamento do mês.',
       '3) Células amarelas são para preencher; as laranjas são o resultado.',
       '',
       'Os números preenchidos são só exemplos. Nada aqui é recomendação contábil: fale com o seu contador sobre custos e impostos.',
       'Quer ver quanto sobrou no mês sem planilha? Teste 14 dias grátis, sem cartão: https://tanocaixa.com']
for i, t in enumerate(txt, start=1): w3.cell(i, 1, t).font = Font(name='Arial', bold=(i == 1), size=14 if i == 1 else 11)
w3.column_dimensions['A'].width = 120

wb.save('public/downloads/planilha-cmv-ficha-tecnica-tanocaixa.xlsx')
print('ok')
